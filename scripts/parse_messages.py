#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""微信聊天导出文件（WeFlow 格式）的解析工具。

可作为模块 import，也可直接跑做结构自检：
    python parse_messages.py <聊天导出.xlsx>

依赖: openpyxl
"""
import collections
import re
import sys
from datetime import date, datetime

import openpyxl

FW = str.maketrans('０１２３４５６７８９', '0123456789')
PUNCT = '（）()【】[]。，,.、:：;；"\'“”‘’!！?？·~—_-'
MOBILE = re.compile(r'1[3-9]\d{9}')
CJK = re.compile(r'[\u4e00-\u9fff]')
MARK = r'(?:微信|微|号|同号)'
ADDR_KW = re.compile(r'省|市|区|县|镇|街道|路|村|号|组|栋|单元|室|收|电话|手机')
BADGE = re.compile(r'顺丰|顺心捷达|京东|包邮|到付|陆运|快递|代收|发货|单号|SF\d+')
PRICE = re.compile(r'[（(]\s*\d+\s*元?\s*[)）]')
PROV_FULL = re.compile(
    r'((?:北京|天津|上海|重庆)市|[\u4e00-\u9fff]{2,3}省|内蒙古自治区|新疆维吾尔自治区'
    r'|广西壮族自治区|宁夏回族自治区|西藏自治区)')
PROV_SHORT = re.compile(
    r'(河北|山西|辽宁|吉林|黑龙江|江苏|浙江|安徽|福建|江西|山东|河南|湖北|湖南|广东|海南'
    r'|四川|贵州|云南|陕西|甘肃|青海|台湾|内蒙古|广西|西藏|宁夏|新疆)')
CITY = re.compile(r'([\u4e00-\u9fff]{2,4}市)')
ZHOU = re.compile(r'([\u4e00-\u9fff]{2,6}(?:自治州|地区|盟|州))')
MUNI = ('北京', '天津', '上海', '重庆')

# 省份统一成标准全称（台账里统一用全称，如「广东省」「宁夏回族自治区」）
PROV_NORM = {p: p + '省' for p in (
    '河北 山西 辽宁 吉林 黑龙江 江苏 浙江 安徽 福建 江西 山东 河南 湖北 湖南 '
    '广东 海南 四川 贵州 云南 陕西 甘肃 青海 台湾').split()}
PROV_NORM.update({
    '内蒙古': '内蒙古自治区', '广西': '广西壮族自治区', '西藏': '西藏自治区',
    '宁夏': '宁夏回族自治区', '新疆': '新疆维吾尔自治区',
})


# ---------------------------------------------------------------- 读源文件
def load_chat(path, header_rows=4):
    """读 WeFlow 导出的聊天 xlsx，返回 (meta, messages)。

    meta: {'wxid':..., 'nick':...}
    messages: [{'seq','time','sender','type','content'}]，已跳过前 header_rows 行会话头
    """
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    meta = {}
    if len(rows) >= 2 and rows[1] and rows[1][0] == '微信ID':
        meta['wxid'] = rows[1][1]
        meta['nick'] = rows[1][4] if len(rows[1]) > 4 else None
    msgs = []
    for r in rows[header_rows:]:
        if r is None or r[4] is None:
            continue
        msgs.append(dict(seq=r[0], time=str(r[1]), sender=r[2], type=r[3], content=str(r[4])))
    return meta, msgs


def to_date(t):
    return datetime.strptime(str(t)[:10], '%Y-%m-%d').date()


# ---------------------------------------------------------------- 号码处理
def norm(s):
    s = str(s).translate(FW)
    s = re.sub(r'[\s\u3000]', '', s)
    return s.strip(PUNCT)


def classify(raw):
    """把一条消息 / 一个 token 归成 ('phone'|'wechat', 值, 备注) 或 None。

    判定顺序不可调换，理由见 ../references/number-formats.md
    """
    t = norm(raw)
    if not t:
        return None
    if re.fullmatch(r'1[3-9]\d{9}', t):
        return ('phone', t, '')
    if t.isdigit():
        return ('phone', t, '仅 %d 位，非 11 位手机号，待核对' % len(t))

    t2 = re.sub(r'^%s+' % MARK, '', t)
    t2 = re.sub(r'%s+$' % MARK, '', t2).strip(PUNCT)
    if not t2:
        return None
    if not MOBILE.search(t2) and CJK.search(t2):        # 正文
        return None
    if re.search(r'[A-Za-z]', t2):                      # 微信号
        note = '原文「%s」' % raw if t2 != t else ''
        if t2[:1] in 'Qq' and t2[1:].isdigit():
            note = '原文「%s」；Q 开头且为纯数字，疑为 QQ 号，待确认' % raw
        return ('wechat', t2, note)
    m = MOBILE.search(t2)
    if m:
        extra = (t2[:m.start()] + t2[m.end():]).strip(PUNCT)
        if extra:
            note = '原文「%s」；修饰词「%s」已剔除' % (raw, extra)
        elif t2 != t:
            note = '原文「%s」；中文标记已剔除' % raw
        else:
            note = ''
        return ('phone', m.group(0), note)
    return None


def find_mobiles(text):
    return MOBILE.findall(str(text))


def is_bare_number(text):
    """整条消息就是一个号码（可能带 v、空格、括号）。"""
    s = norm(text)
    return bool(re.fullmatch(r'[vV]?\d{10,16}[vV]?', s))


def looks_like_order(text):
    """群聊发货记录的判据：同时含号码和地址关键词。"""
    return bool(MOBILE.search(str(text))) and bool(ADDR_KW.search(str(text)))


# ---------------------------------------------------------------- 地区
def parse_region(text):
    """从一段地址文本抽 (省份, 城市)。只认原文写出的，不反推。"""
    t = str(text)
    prov = city = ''
    m = PROV_FULL.search(t)
    if not m:
        m = PROV_SHORT.search(t)
    if m:
        raw = m.group(1)
        base = (raw.replace('维吾尔自治区', '').replace('壮族自治区', '')
                .replace('回族自治区', '').replace('自治区', '').replace('省', ''))
        if base in MUNI:
            return base + '市', ''          # 直辖市：省份填市名，城市列留空
        prov = PROV_NORM.get(base, base)    # 统一成标准全称：云南 -> 云南省
        t = t.replace(raw, '', 1)
    # 先把「XX自治州/XX地区/XX盟」整段摘掉再找市，
    # 否则 '西双版纳傣族自治州景洪市' 会被贪婪匹配切成「治州景洪市」
    t_no_zhou = ZHOU.sub('', t)
    city = _first_city(t_no_zhou)
    if not city:
        m = ZHOU.search(t)
        if m:
            city = m.group(1)
    return prov, city


BAD_CITY_CHARS = ('省', '区', '县', '族', '治')   # 注意：「州」不能排除，杭州/苏州都含州


def _first_city(text):
    """取第一个「XX市」，但候选里不能混进 省/区/县/族/治 这类字。

    排除「治」是为了挡掉 '自治州景洪市' 被贪婪切成「治州景洪市」的情况，
    同时又不能误伤「杭州市」「苏州市」。
    """
    for m in CITY.finditer(text):
        body = m.group(1)[:-1]
        if any(ch in body for ch in BAD_CITY_CHARS):
            continue
        return m.group(1)
    return ''


# ---------------------------------------------------------------- 汇总
def collect_contacts(messages, sender=None, dedupe=True):
    """扫全部消息，抽出 (kind, value) 去重，日期取最早。

    sender: 只看某个发送者（私聊传 '我'；群聊传 None）
    """
    seen = {}
    for msg in messages:
        if sender is not None and msg['sender'] != sender:
            continue
        res = classify(msg['content'])
        if res is None:
            continue
        kind, val, note = res
        d = to_date(msg['time'])
        key = (kind, val)
        if key not in seen or d < seen[key]['d']:
            seen[key] = dict(kind=kind, val=val, note=note, d=d)
    out = list(seen.values())
    return sorted(out, key=lambda x: (x['d'], x['val'])) if dedupe else out


def strip_device(text):
    """从发货指令里剥出设备描述（去掉物流词、价格、地址行）。"""
    keep = []
    for line in re.split(r'[\n，,。;；]', str(text)):
        s = norm(line)
        if not s:
            continue
        s = MOBILE.sub('', s)
        s = BADGE.sub('', s)
        s = PRICE.sub('', s).strip('。，,、()（）')
        if not s or len(s) > 60:
            continue
        if ADDR_KW.search(s):
            continue
        keep.append(s)
    return '；'.join(keep)


# ---------------------------------------------------------------- 校验
def verify(xlsx_path, sheet='客户信息'):
    """回读台账并打印校验结果。任何一项异常都说明写入环节踩坑了。"""
    ws = openpyxl.load_workbook(xlsx_path)[sheet]
    last = 0
    for r in range(2, 5000):
        if any(ws.cell(row=r, column=c).value for c in (2, 3, 4)):
            last = r
    n = last - 1 if last else 0
    rows = [tuple(ws.cell(row=r, column=c).value for c in range(2, 11)) for r in range(2, last + 1)]
    nums = []
    for x in rows:
        if x[1]:
            nums += [t for t in str(x[1]).split('、') if t]
        if x[2]:
            nums.append(x[2])
    dup = [k for k, v in collections.Counter(nums).items() if v > 1]
    names = ['客户昵称', '手机号', '微信号', '聊天日期', '发货日期', '省份', '城市', '设备', '备注']
    print('文件:', xlsx_path)
    print('数据行数:', n, '(空模板，尚未填数据)' if n == 0 else '')
    print('重复号码:', dup or '无')
    print('缺号码的行:', sum(1 for x in rows if not x[1] and not x[2]))
    for i, h in enumerate(names):
        print('  %-6s %d' % (h, sum(1 for x in rows if x[i])))
    ok = not dup                      # 空模板是合法状态，不算异常
    print('结论:', 'OK' if ok else '!! 有重复号码，需检查')
    return ok


def overlap(paths, sheet='客户信息'):
    """跨表号码交叉比对。"""
    def keyset(p):
        ws = openpyxl.load_workbook(p)[sheet]
        s = set()
        for r in range(2, 5000):
            v3, v4 = ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value
            if v3:
                s |= {t for t in str(v3).split('、') if t}
            if v4:
                s.add(str(v4))
        return s
    sets = {p: keyset(p) for p in paths}
    names = list(sets)
    for i in range(len(names)):
        for j in range(i + 1, len(names)):
            ov = sorted(sets[names[i]] & sets[names[j]])
            print('%s ∩ %s: %s' % (names[i], names[j], ov or '无'))
    allk = set().union(*sets.values())
    print('合计去重号码:', len(allk))
    return sets


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    path = sys.argv[1]
    meta, msgs = load_chat(path)
    print('会话:', meta)
    print('消息数:', len(msgs))
    print('时间范围:', msgs[0]['time'], '~', msgs[-1]['time'])
    print('发送者分布:', collections.Counter(m['sender'] for m in msgs).most_common(10))
    print('消息类型分布:', collections.Counter(m['type'] for m in msgs).most_common())
    orders = [m for m in msgs if looks_like_order(m['content'])]
    print('含号码+地址的消息（疑似发货记录）:', len(orders))
    print()
    print('--- 前 5 行原文预览 ---')
    for r in msgs[:5]:
        print(' ', r['time'], '|', r['sender'], '|', r['content'][:80].replace('\n', ' '))
    print()
    print('提示：先看这个自检输出确认列结构，再决定走「批量号码」还是「人工校对」路线。')
    return 0


if __name__ == '__main__':
    sys.exit(main())
