#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成「客户信息台账」空白模板。

用法:
    python build_table.py <输出.xlsx> [标题] [sheet名逗号分隔]

示例:
    python build_table.py ./客户信息收集表.xlsx "客户信息收集表"
    python build_table.py ./表_发货群.xlsx "客户信息收集表_发货群A"

不传 sheet 名时默认：客户信息,填写概览,填写说明
字段规格见 ../references/field-spec.md

注意: 本脚本每次都新建工作簿（从零生成），不会在既有文件上叠加——
这是为了规避 openpyxl 的 ws.cell(row, col, value=None) 不清空单元格的坑。
"""
import os
import sys

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LAST = 1000
BLUE, LIGHT = "4472C4", "D9E2F3"
YELLOW_BG, YELLOW_FG, WHITE = "FFEB9C", "9C6500", "FFFFFF"

HEADERS = ["序号", "客户昵称", "手机号", "微信号", "聊天日期", "发货日期",
           "省份", "城市", "购买设备", "备注"]
TEXT_COLS = [3, 4, 7, 8, 9, 10]      # 强制文本，防号码变科学计数法
DATE_COLS = [5, 6]
WIDTHS = {"A": 8, "B": 16, "C": 17, "D": 18, "E": 13, "F": 13,
          "G": 12, "H": 12, "I": 34, "J": 40}

thin_bottom = Border(bottom=Side(style="thin", color="BFBFBF"))
head_font = Font(name="等线", size=11, bold=True, color=WHITE)
head_fill = PatternFill("solid", fgColor=BLUE)
head_align = Alignment(horizontal="center", vertical="center")
body_font = Font(name="等线", size=11)
wrap_top = Alignment(vertical="top", wrap_text=True)

NOTES = [
    ("客户昵称", "微信昵称，或你自己给客户起的备注名，用来认人。同一客户多次出现只占一行。"),
    ("手机号", "优先从发货记录里取。没有的留空，不要猜、不要拿微信号凑数。多个号码写在同一格，用「、」分隔。"),
    ("微信号", "聊天记录中出现的微信号 / 微信 ID。只有昵称、没有明确微信号的留空。"),
    ("聊天日期", "该客户这条聊天记录发生的日期，格式 2026-09-15。多条聊天取最早一条作首次接触日期。"),
    ("发货日期", "发货记录 / 订单上的日期，格式 2026-09-15。没发货或记录里没有的留空。"),
    ("省份 / 城市", "只填原文写出的，不按城市反推省份。原文只写到县、或没写「市」字的，城市列留空。"),
    ("购买设备", "沿用客户自己的叫法。多台写在同一格，用「、」分隔，例：扫地机×2、洗地机×1。"),
    ("备注", "原文地址、信息来源、以及拿不准需要回头确认的地方。建议以「原文地址：」开头。"),
    ("行口径", "一行 = 一个客户。同一客户买了多台设备、分多次下单，都合并到同一行。"),
]


def build(path, title, sheets):
    s_info, s_kpi, s_note = sheets
    wb = openpyxl.Workbook()
    wb.properties.title = title
    wb.properties.creator = title

    # ---------- 主表 ----------
    ws = wb.active
    ws.title = s_info
    ncol = len(HEADERS)
    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=i)
        c.value = h
        c.font, c.fill, c.alignment, c.border = head_font, head_fill, head_align, thin_bottom
    ws.row_dimensions[1].height = 26
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    for r in range(2, LAST + 1):
        ws.cell(row=r, column=1).value = '=IF($B{r}="","",ROW()-1)'.format(r=r)
        for col in range(1, ncol + 1):
            cell = ws.cell(row=r, column=col)
            cell.font, cell.alignment = body_font, wrap_top
            if col in TEXT_COLS:
                cell.number_format = "@"
            elif col in DATE_COLS:
                cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:{0}{1}".format(get_column_letter(ncol), LAST)
    ws.conditional_formatting.add(
        "C2:D{0}".format(LAST),
        FormulaRule(formula=['AND($B2<>"",$C2="",$D2="")'],
                    fill=PatternFill("solid", bgColor=YELLOW_BG),
                    font=Font(color=YELLOW_FG), stopIfTrue=False),
    )

    # ---------- 填写概览 ----------
    w2 = wb.create_sheet(s_kpi)
    w2.merge_cells("A1:B1")
    t = w2["A1"]
    t.value, t.font, t.fill, t.alignment = (title + " · 填写情况概览"), Font(name="等线", size=12, bold=True, color=WHITE), head_fill, head_align
    w2.row_dimensions[1].height = 26
    for i, h in enumerate(["统计项", "数值"], start=1):
        c = w2.cell(row=2, column=i)
        c.value = h
        c.font = Font(name="等线", size=11, bold=True)
        c.fill = PatternFill("solid", fgColor=LIGHT)
        c.alignment, c.border = head_align, thin_bottom
    S = s_info + "!"

    def cnt(letter):
        return '=SUMPRODUCT(--({0}${1}$2:${1}${2}<>""))'.format(S, letter, LAST)

    kpis = [("客户总数", "=COUNTA({0}$B$2:$B${1})".format(S, LAST)),
            ("手机号已填", cnt("C")), ("微信号已填", cnt("D")),
            ("手机号缺失", "=$B$3-$B$4"), ("微信号缺失", "=$B$3-$B$5"),
            ("聊天日期已填", cnt("E")), ("发货日期已填", cnt("F")),
            ("省份已填", cnt("G")), ("城市已填", cnt("H")), ("设备已填", cnt("I"))]
    for idx, (label, formula) in enumerate(kpis):
        r = 3 + idx
        a = w2.cell(row=r, column=1)
        a.value, a.font, a.alignment = label, body_font, Alignment(vertical="center")
        b = w2.cell(row=r, column=2)
        b.value = formula
        b.font = Font(name="等线", size=11, bold=True)
        b.alignment = Alignment(horizontal="center", vertical="center")
        b.number_format = "#,##0"
    w2.column_dimensions["A"].width, w2.column_dimensions["B"].width = 23, 12

    # ---------- 填写说明 ----------
    w3 = wb.create_sheet(s_note)
    w3.merge_cells("A1:B1")
    t3 = w3["A1"]
    t3.value = "字段填写规则"
    t3.font, t3.fill, t3.alignment = Font(name="等线", size=12, bold=True, color=WHITE), head_fill, head_align
    w3.row_dimensions[1].height = 26
    for i, h in enumerate(["字段", "填写规则"], start=1):
        c = w3.cell(row=2, column=i)
        c.value = h
        c.font = Font(name="等线", size=11, bold=True)
        c.fill = PatternFill("solid", fgColor=LIGHT)
        c.alignment, c.border = head_align, thin_bottom
    for idx, (f, rule) in enumerate(NOTES):
        r = 3 + idx
        a = w3.cell(row=r, column=1)
        a.value, a.font, a.alignment = f, Font(name="等线", size=11, bold=True), wrap_top
        b = w3.cell(row=r, column=2)
        b.value, b.font, b.alignment = rule, body_font, wrap_top
    w3.column_dimensions["A"].width, w3.column_dimensions["B"].width = 23, 60

    wb.save(path)
    return path


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    path = os.path.abspath(sys.argv[1])
    title = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(os.path.basename(path))[0]
    sheets = [s for s in (sys.argv[3].split(",") if len(sys.argv) > 3
                          else ["客户信息", "填写概览", "填写说明"]) if s]
    if len(sheets) != 3:
        print("!! 需要 3 个 sheet 名", file=sys.stderr)
        return 2
    build(path, title, sheets)
    print("OK ->", path)
    print("sheets:", sheets)
    return 0


if __name__ == "__main__":
    sys.exit(main())
